"""Импорт списка товаров из файлов (.xlsx, .csv, .txt) и экспорт результатов в Excel."""
import csv
import io
import math
import re
from collections.abc import Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from app.schemas import REGIME_LABELS, ClassificationResult

HEADER_WORDS = re.compile(r"^(№|n|наименование|название|товар|продукция|позиция|предмет|name|item)", re.IGNORECASE)
_has_letters = re.compile(r"[a-zа-яё]{2,}", re.IGNORECASE)


class FileReadError(Exception):
    pass


def _looks_like_name(value) -> bool:
    return isinstance(value, str) and len(value.strip()) >= 3 and bool(_has_letters.search(value))


def _pick_column(rows: list[list]) -> list[str]:
    """Колонка, где больше всего «наименований» (текст с буквами), без строки-шапки."""
    if not rows:
        return []
    width = max(len(r) for r in rows)
    best, best_score = 0, -1.0
    for col in range(width):
        values = [r[col] for r in rows if col < len(r) and r[col] not in (None, "")]
        if not values:
            continue
        names = sum(_looks_like_name(v) for v in values)
        avg_len = sum(len(str(v)) for v in values) / len(values)
        score = names / len(values) + min(avg_len, 60) / 600  # при равенстве — колонка с более длинным текстом
        if names / len(values) >= 0.6 and score > best_score:
            best, best_score = col, score
    out = [str(r[best]).strip() for r in rows if best < len(r) and _looks_like_name(r[best])]
    if out and HEADER_WORDS.match(out[0]):
        out = out[1:]
    return out


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise FileReadError("Не удалось определить кодировку файла (ожидается UTF-8 или Windows-1251)")


def read_names(filename: str, data: bytes) -> list[str]:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "xlsx":
        try:
            ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
        except Exception:
            raise FileReadError("Не удалось открыть Excel-файл. Сохраните его в формате .xlsx и попробуйте снова")
        return _pick_column(rows)
    if ext == "csv":
        text = _decode(data)
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";" if text.count(";") > text.count(",") else ","
        return _pick_column([row for row in csv.reader(io.StringIO(text), dialect)])
    if ext == "txt":
        return [line.strip() for line in _decode(data).splitlines() if _has_letters.search(line)]
    raise FileReadError("Неподдерживаемый формат файла: загрузите .xlsx, .csv или .txt")


EXPORT_COLUMNS = ["Товар", "Код ОКПД2", "Наименование ОКПД2", "Режим ПП 1875", "Пункты ПП 1875",
                  "Уверенность, %", "Проверить вручную", "Пояснение"]
REGIME_FILL = {"ban": "F8D7DA", "restriction": "FFF3CD", "advantage": "D6E4FF", "none": "D4EDDA"}


def results_to_xlsx(results: Sequence[ClassificationResult]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Результаты"
    ws.append(EXPORT_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in results:
        items = "; ".join(
            f"прил. {m.appendix} п. {m.item_no}" + (f" (доля {m.quota}%)" if m.quota else "") for m in r.npa_matches
        )
        ws.append([
            r.input,
            r.okpd2.code if r.okpd2 else "",
            r.okpd2.title if r.okpd2 else "",
            REGIME_LABELS[r.verdict],
            items or "—",
            math.floor(r.confidence * 100 + 1e-9),  # вниз, как на сайте: 0.7964 при пороге 80% — это 79%
            "да" if r.needs_review else "нет",
            r.error or r.explanation,
        ])
        ws.cell(ws.max_row, 4).fill = PatternFill("solid", fgColor=REGIME_FILL[r.verdict.value])
    for col, width in zip("ABCDEFGH", (40, 16, 50, 16, 30, 14, 12, 60)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
