"""Разбор файла редакции ПП № 1875 (JSON или XLSX) в список позиций."""
import json
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from app.schemas import REGIME_BY_APPENDIX, Regime

CODE_RE = re.compile(r"\d{2}(?:\.\d{1,3}){0,3}")
CONDITIONS_RE = re.compile(r"за исключением|кроме|только|соответствующ", re.IGNORECASE)
XLSX_COLUMNS = ["appendix", "item_no", "title", "codes", "quota", "exclusions", "note"]


class NpaParseError(Exception):
    pass


@dataclass
class NpaItemIn:
    appendix: int
    item_no: str
    title: str
    codes: list[str]
    regime: Regime
    quota: int | None = None
    exclusions: list[str] = field(default_factory=list)
    note: str = ""
    has_conditions: bool = False


@dataclass
class NpaDocument:
    title: str
    source_url: str
    items: list[NpaItemIn]


def _codes(value) -> list[str]:
    if value is None:
        return []
    if isinstance(value, (list, tuple)):
        value = ",".join(str(v) for v in value)
    return CODE_RE.findall(str(value).replace("\xa0", " "))


def _item(raw: dict, where: str) -> NpaItemIn:
    try:
        appendix = int(raw["appendix"])
        regime = REGIME_BY_APPENDIX[appendix]
    except (KeyError, TypeError, ValueError):
        raise NpaParseError(f"{where}: номер приложения должен быть 1, 2 или 3")
    title = str(raw.get("title") or "").strip()
    codes = _codes(raw.get("codes"))
    if not codes:
        raise NpaParseError(f"{where}: не указан ни один код ОКПД2")
    quota = raw.get("quota")
    return NpaItemIn(
        appendix=appendix,
        item_no=str(raw.get("item_no") or "").strip(),
        title=title,
        codes=codes,
        regime=regime,
        quota=int(quota) if quota not in (None, "") else None,
        exclusions=_codes(raw.get("exclusions")),
        note=str(raw.get("note") or "").strip(),
        has_conditions=bool(CONDITIONS_RE.search(title)),
    )


def _parse_json(path: Path) -> NpaDocument:
    try:
        doc = json.loads(path.read_text(encoding="utf-8-sig"))
    except (json.JSONDecodeError, UnicodeDecodeError) as e:
        raise NpaParseError(f"Не удалось прочитать JSON: {e}")
    raw_items = doc.get("items") if isinstance(doc, dict) else None
    if not isinstance(raw_items, list):
        raise NpaParseError("Не удалось найти список позиций (поле items)")
    items = [_item(r, f"Позиция {n}") for n, r in enumerate(raw_items, 1)]
    return NpaDocument(title=str(doc.get("title") or ""), source_url=str(doc.get("source_url") or ""), items=items)


def _parse_xlsx(path: Path) -> NpaDocument:
    try:
        ws = load_workbook(path, read_only=True, data_only=True).active
    except Exception as e:
        raise NpaParseError(f"Не удалось открыть Excel-файл: {e}")
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip().lower() for h in next(rows, [])]
    missing = [c for c in ("appendix", "item_no", "codes") if c not in header]
    if missing:
        raise NpaParseError(f"В первой строке нет колонок: {', '.join(missing)}. Ожидаются: {', '.join(XLSX_COLUMNS)}")
    items = []
    for n, row in enumerate(rows, 2):
        if not row or all(v in (None, "") for v in row):
            continue
        items.append(_item(dict(zip(header, row)), f"Строка {n}"))
    return NpaDocument(title="", source_url="", items=items)


def parse_npa(path: Path) -> NpaDocument:
    suffix = path.suffix.lower()
    if suffix == ".json":
        doc = _parse_json(path)
    elif suffix == ".xlsx":
        doc = _parse_xlsx(path)
    else:
        raise NpaParseError("Неподдерживаемый формат файла: нужен .json или .xlsx")
    if not doc.items:
        raise NpaParseError("В файле не найдено ни одной позиции")
    return doc
