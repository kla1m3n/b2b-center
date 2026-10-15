from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.io_files import FileReadError, read_names, results_to_xlsx
from app.schemas import ClassificationResult, NpaMatch, Okpd2Ref, Regime


def xlsx_bytes(rows) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_read_names_excel_with_header_numbers_and_empty_rows():
    data = xlsx_bytes([
        ["№", "Наименование товара", "Кол-во"],
        [1, "Ноутбук Lenovo IdeaPad 15", 10],
        [None, None, None],
        [2, "Бумага А4", 500],
        [3, "Стул офисный", 20],
    ])
    assert read_names("spec.xlsx", data) == ["Ноутбук Lenovo IdeaPad 15", "Бумага А4", "Стул офисный"]


def test_read_names_cp1251_semicolon_csv():
    data = "Наименование;Количество\nБумага офисная А4;10\nРучка шариковая;5\n".encode("cp1251")
    assert read_names("a.csv", data) == ["Бумага офисная А4", "Ручка шариковая"]


def test_read_names_utf8_comma_csv_without_header():
    data = "Ноутбук,1\nМонитор 27,2\n".encode("utf-8-sig")
    assert read_names("b.csv", data) == ["Ноутбук", "Монитор 27"]


def test_read_names_plain_text_lines():
    assert read_names("list.txt", "Ноутбук\n\n  Стул  \n".encode()) == ["Ноутбук", "Стул"]


def test_unsupported_or_broken_file_has_russian_error():
    with pytest.raises(FileReadError, match="формат"):
        read_names("a.pdf", b"%PDF")
    with pytest.raises(FileReadError, match="Не удалось"):
        read_names("a.xlsx", b"not excel")


def test_results_to_xlsx_has_russian_columns_and_rows():
    r = ClassificationResult(
        input="Ноутбук", okpd2=Okpd2Ref(code="26.20.11.110", title="Ноутбуки"), confidence=0.92, needs_review=False,
        verdict=Regime.restriction,
        npa_matches=[NpaMatch(regime=Regime.restriction, appendix=2, item_no="195", title="Компьютеры",
                              matched_code="26.20.11")],
    )
    ws = load_workbook(BytesIO(results_to_xlsx([r]))).active
    header = [c.value for c in ws[1]]
    assert header[:4] == ["Товар", "Код ОКПД2", "Наименование ОКПД2", "Режим ПП 1875"]
    row = [c.value for c in ws[2]]
    assert row[0] == "Ноутбук" and row[1] == "26.20.11.110" and row[3] == "ОГРАНИЧЕНИЕ"
    assert "прил. 2 п. 195" in row[4] and row[5] == 92
